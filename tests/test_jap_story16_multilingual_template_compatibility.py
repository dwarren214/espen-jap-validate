"""Story 16 tests for multilingual JRSM compatibility and false-positive reduction."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from espen_sql_api.jap_validation.validators.jrsm import validate_jrsm_workbook

REQUIRED_SHEETS = [
    "INTRO",
    "COUNTRY_INFO",
    "DEC",
    "ALB_MBD",
    "PZQ",
    "IVM",
    "IVM+",
    "SUMMARY",
    "SHIPMENT",
    "DATA_POLICY",
]
SUPPORTED_LANGUAGE_TEMPLATES = {
    "en": Path("docs/JRSM-template-all-sheets.xlsx"),
    "fr": Path("docs/WHO_JRSM_PC_v4_fr (1).xlsm"),
    "es": Path("docs/WHO_JRSM_PC_v4_es.xlsm"),
}


def _build_minimal_workbook(tmp_path: Path, *, marker: str, e37: str, e39: str, e41: str, e43: str) -> Path:
    workbook = Workbook()
    intro_sheet = workbook.active
    intro_sheet.title = "INTRO"

    for sheet_name in REQUIRED_SHEETS:
        if sheet_name == "INTRO":
            continue
        workbook.create_sheet(title=sheet_name)

    intro_sheet["B2"] = marker
    intro_sheet["E33"] = "Rwanda"
    intro_sheet["E35"] = 2026
    intro_sheet["E37"] = e37
    intro_sheet["E39"] = e39
    intro_sheet["E41"] = e41
    intro_sheet["E43"] = e43
    intro_sheet["E45"] = 1
    intro_sheet["E48"] = 0.15
    intro_sheet["E49"] = 0.25
    intro_sheet["E50"] = 0.6
    intro_sheet["E51"] = 0.3

    workbook["COUNTRY_INFO"]["A10"] = "Rwanda"
    workbook["COUNTRY_INFO"]["B10"] = "Province"
    workbook["COUNTRY_INFO"]["C10"] = "District"
    workbook["COUNTRY_INFO"]["D10"] = 100
    workbook["COUNTRY_INFO"]["H10"] = 1
    workbook["COUNTRY_INFO"]["I10"] = 1
    workbook["COUNTRY_INFO"]["J10"] = 2
    workbook["COUNTRY_INFO"]["K10"] = 2
    workbook["COUNTRY_INFO"]["V10"] = 1
    workbook["COUNTRY_INFO"]["W10"] = 1
    workbook["COUNTRY_INFO"]["X10"] = 1
    workbook["COUNTRY_INFO"]["Y10"] = 1
    workbook["COUNTRY_INFO"]["Z10"] = "IU"
    workbook["COUNTRY_INFO"]["AB10"] = 1
    workbook["COUNTRY_INFO"]["AC10"] = 1
    workbook["COUNTRY_INFO"]["AD10"] = 1

    path = tmp_path / "story16.xlsx"
    workbook.save(path)
    return path


def test_blank_but_addressable_intro_spacer_cells_do_not_emit_required_cell_findings(tmp_path: Path):
    workbook_path = _build_minimal_workbook(
        tmp_path,
        marker="Joint request for selected PC medicines v.4.4",
        e37="Endemic",
        e39="Endemic",
        e41="Endemic",
        e43="Endemic",
    )

    result = validate_jrsm_workbook(
        workbook_path=workbook_path,
        country="Rwanda",
        year_for_request_of_medicine=2026,
    )

    required_cell_findings = [finding for finding in result.findings if finding.rule_id == "JRSM.TEMPLATE.REQUIRED_CELLS"]
    assert not required_cell_findings


def test_supported_templates_emit_language_profile_and_avoid_formula_mismatch_noise():
    for language, workbook_path in SUPPORTED_LANGUAGE_TEMPLATES.items():
        result = validate_jrsm_workbook(
            workbook_path=workbook_path,
            country="Rwanda",
            year_for_request_of_medicine=2026,
        )

        language_profile_findings = [
            finding for finding in result.findings if finding.rule_id == "JRSM.TEMPLATE.LANGUAGE_PROFILE"
        ]
        assert language_profile_findings
        assert any(finding.severity == "info" and finding.actual == language for finding in language_profile_findings)

        version_findings = [finding for finding in result.findings if finding.rule_id == "JRSM.TEMPLATE.VERSION_MARKER"]
        assert version_findings
        assert any(finding.severity == "info" for finding in version_findings)

        formula_match_findings = [
            finding for finding in result.findings if finding.rule_id == "JRSM.FORMULA.GOVERNED_FORMULA_MATCH"
        ]
        assert not formula_match_findings


def test_legacy_french_marker_skips_formula_governance_noise(tmp_path: Path):
    workbook_path = _build_minimal_workbook(
        tmp_path,
        marker="Formulaire de demande commune de médicaments pour CP v.3.1 (L)",
        e37="Endémique",
        e39="Endémique",
        e41="Endémique",
        e43="Endémique",
    )

    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=False)
    workbook["COUNTRY_INFO"]["N10"] = (
        '=IF(INTRO!$E$41="Endémique", IF($J10>1, IF($J10=4, "Inconnu", '
        'IF($J10=99, "Surveillance", IF($J10=11,E10*0.5, E10))), 0),"Non requis")'
    )
    workbook["IVM"]["D10"] = (
        '=IF(INTRO!$E$39="Non endémique",0,IF(INTRO!$E$39="Endémique mais pas de CP",'
        'IF(COUNTRY_INFO!$H10=1,COUNTRY_INFO!F10+COUNTRY_INFO!G10,0),'
        'IF(COUNTRY_INFO!$H10=1,COUNTRY_INFO!F10+COUNTRY_INFO!G10,'
        'IF(COUNTRY_INFO!$I10=1,IF(COUNTRY_INFO!M10>=(COUNTRY_INFO!F10+COUNTRY_INFO!G10),'
        "COUNTRY_INFO!F10+COUNTRY_INFO!G10,COUNTRY_INFO!$M10),0))))"
    )
    workbook.save(workbook_path)

    result = validate_jrsm_workbook(
        workbook_path=workbook_path,
        country="Rwanda",
        year_for_request_of_medicine=2026,
    )

    language_profile_findings = [
        finding for finding in result.findings if finding.rule_id == "JRSM.TEMPLATE.LANGUAGE_PROFILE"
    ]
    assert language_profile_findings
    assert any(finding.severity == "warn" and finding.actual.endswith("v.3.1 (L)") for finding in language_profile_findings)

    version_findings = [finding for finding in result.findings if finding.rule_id == "JRSM.TEMPLATE.VERSION_MARKER"]
    assert version_findings
    assert any(finding.severity == "warn" for finding in version_findings)

    formula_findings = [finding for finding in result.findings if finding.rule_id.startswith("JRSM.FORMULA.")]
    assert not formula_findings
