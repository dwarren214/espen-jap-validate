"""Story 18 tests for tolerant JRSM workbook parsing."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import zipfile

from openpyxl import Workbook

from espen_sql_api.jap_validation.validators import jrsm
from espen_sql_api.jap_validation.validators.jrsm import validate_jrsm_workbook


FRENCH_V43_MARKER = "Formulaire de demande commune de médicaments pour CP v.4.3"


def _build_minimal_jrsm_workbook(path: Path) -> None:
    workbook = Workbook()
    intro_sheet = workbook.active
    intro_sheet.title = "INTRO"

    for sheet_name in jrsm.REQUIRED_SHEETS:
        if sheet_name != "INTRO":
            workbook.create_sheet(title=sheet_name)

    for sheet in workbook.worksheets:
        sheet.column_dimensions["A"].width = 12.5

    intro_sheet["B2"] = FRENCH_V43_MARKER
    intro_sheet["E33"] = "BENIN"
    intro_sheet["E35"] = 2026
    intro_sheet["E37"] = "Endémique mais pas de CP"
    intro_sheet["E39"] = "Endémique"
    intro_sheet["E41"] = "Endémique"
    intro_sheet["E43"] = "Endémique"
    intro_sheet["E45"] = 1
    intro_sheet["E48"] = 0.15
    intro_sheet["E49"] = 0.25
    intro_sheet["E50"] = 0.6
    intro_sheet["E51"] = 0.3

    country_info = workbook["COUNTRY_INFO"]
    country_info["H7"] = "FL"
    country_info["I7"] = "Oncho"
    country_info["J7"] = "STH"
    country_info["K7"] = "SCH"
    country_info["B10"] = "Province"
    country_info["C10"] = "District"
    country_info["D10"] = "Commune"
    country_info["H10"] = 1
    country_info["I10"] = 1
    country_info["J10"] = 1
    country_info["K10"] = 1

    workbook["SUMMARY"]["C6"] = '=IF(INTRO!$E$33<>0,INTRO!$E$33," ")'
    workbook["SHIPMENT"]["C4"] = '=IF(INTRO!$E$33<>0,INTRO!$E$33," ")'
    workbook["DEC"].sheet_state = "hidden"
    workbook["IVM+"].sheet_state = "hidden"
    workbook["DATA_POLICY"].sheet_state = "hidden"
    workbook.save(path)


def _inject_widthpt_into_first_col_tag(xml_bytes: bytes) -> tuple[bytes, int]:
    start = xml_bytes.find(b"<col ")
    if start < 0:
        return xml_bytes, 0

    end = xml_bytes.find(b">", start)
    if end < 0:
        return xml_bytes, 0

    tag = xml_bytes[start:end]
    if b"widthPt=" in tag:
        return xml_bytes, 0

    stripped_tag = tag.rstrip()
    if stripped_tag.endswith(b"/"):
        insert_at = start + len(stripped_tag) - 1
    else:
        insert_at = end
    return xml_bytes[:insert_at] + b' widthPt="7.5"' + xml_bytes[insert_at:], 1


def _add_widthpt_to_worksheet_columns(path: Path) -> int:
    source_payload = path.read_bytes()
    mutated_payload = BytesIO()
    injected_count = 0

    with zipfile.ZipFile(BytesIO(source_payload), "r") as source_zip:
        with zipfile.ZipFile(mutated_payload, "w") as target_zip:
            target_zip.comment = source_zip.comment
            for source_info in source_zip.infolist():
                payload = source_zip.read(source_info.filename)
                if source_info.filename.startswith("xl/worksheets/") and source_info.filename.endswith(".xml"):
                    payload, count = _inject_widthpt_into_first_col_tag(payload)
                    injected_count += count
                target_zip.writestr(source_info, payload)

    path.write_bytes(mutated_payload.getvalue())
    return injected_count


def _build_widthpt_workbook(path: Path) -> Path:
    _build_minimal_jrsm_workbook(path)
    injected_count = _add_widthpt_to_worksheet_columns(path)
    assert injected_count == len(jrsm.REQUIRED_SHEETS)
    return path


def test_widthpt_workbook_runs_downstream_validation_without_modifying_source(tmp_path: Path):
    workbook_path = _build_widthpt_workbook(tmp_path / "benin_widthpt.xlsm")
    original_bytes = workbook_path.read_bytes()

    result = validate_jrsm_workbook(
        workbook_path=workbook_path,
        country="Rwanda",
        year_for_request_of_medicine=2025,
    )

    assert workbook_path.read_bytes() == original_bytes
    assert not [finding for finding in result.findings if finding.rule_id == "JRSM.WORKBOOK.FILE_PARSE"]

    country_findings = [finding for finding in result.findings if finding.rule_id == "JRSM.INTRO.COUNTRY_MATCH"]
    assert country_findings
    assert country_findings[0].actual == "BENIN"

    year_findings = [finding for finding in result.findings if finding.rule_id == "JRSM.INTRO.YEAR_MATCH"]
    assert year_findings
    assert year_findings[0].actual == "2026"

    language_findings = [finding for finding in result.findings if finding.rule_id == "JRSM.TEMPLATE.LANGUAGE_PROFILE"]
    assert language_findings
    assert any(finding.severity == "warn" and finding.actual == FRENCH_V43_MARKER for finding in language_findings)


def test_widthpt_tolerant_loader_preserves_formulas(tmp_path: Path, monkeypatch):
    workbook_path = _build_widthpt_workbook(tmp_path / "formula_widthpt.xlsx")
    original_load_workbook = jrsm.load_workbook

    def load_workbook_with_forced_widthpt_failure(filename, data_only=False):
        if isinstance(filename, Path):
            raise TypeError("ColumnDimension.__init__() got an unexpected keyword argument 'widthPt'")
        return original_load_workbook(filename=filename, data_only=data_only)

    monkeypatch.setattr(jrsm, "load_workbook", load_workbook_with_forced_widthpt_failure)

    workbook, context = jrsm._load_workbook_with_widthpt_tolerance(workbook_path)

    assert context["workbook_parse_mode"] == "openpyxl_widthpt_sanitized"
    assert context["widthpt_sanitized_attribute_count"] == len(jrsm.REQUIRED_SHEETS)
    assert workbook["INTRO"]["E33"].value == "BENIN"
    assert workbook["INTRO"]["E35"].value == 2026
    assert str(workbook["SUMMARY"]["C6"].value).startswith("=")


def test_widthpt_sanitizer_only_strips_worksheet_col_attributes():
    assert jrsm._is_worksheet_xml_path("xl/worksheets/sheet1.xml")
    assert not jrsm._is_worksheet_xml_path("xl/worksheets/_rels/sheet1.xml")

    xml_bytes = (
        b'<worksheet widthPt="root">'
        b"<cols>"
        b'<col min="1" max="1" widthPt="7.5" width="1"/>'
        b"<x:col min=\"2\" widthPt='8' width=\"2\"/>"
        b'<c r="A1" widthPt="cell"/>'
        b"</cols>"
        b"</worksheet>"
    )

    sanitized, stripped_count = jrsm._strip_widthpt_attributes_from_worksheet_cols(xml_bytes)

    assert stripped_count == 2
    assert b'<worksheet widthPt="root">' in sanitized
    assert b'<col min="1" max="1" width="1"/>' in sanitized
    assert b'<x:col min="2" width="2"/>' in sanitized
    assert b'<c r="A1" widthPt="cell"/>' in sanitized


def test_normal_workbook_uses_normal_openpyxl_parse_path(tmp_path: Path):
    workbook_path = tmp_path / "normal.xlsx"
    _build_minimal_jrsm_workbook(workbook_path)

    result = validate_jrsm_workbook(
        workbook_path=workbook_path,
        country="BENIN",
        year_for_request_of_medicine=2026,
    )

    assert result.context["workbook_parse_mode"] == "openpyxl"
    assert result.context["widthpt_sanitized_attribute_count"] == 0
    assert not [finding for finding in result.findings if finding.rule_id == "JRSM.WORKBOOK.FILE_PARSE"]
