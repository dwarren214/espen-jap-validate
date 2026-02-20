"""Story 10 tests for SUMMARY/SHIPMENT placeholder validation."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from espen_sql_api.jap_validation.validators import jrsm

TEMPLATE_WORKBOOK_PATH = Path("docs/JRSM-template-all-sheets.xlsx")
STORY10_RULE_SPEC = "JRSM.SUMMARY_SHIPMENT.SPEC_NOT_CONFIGURED"
STORY10_RULE_ACCESS = "JRSM.SUMMARY_SHIPMENT.BASIC_ACCESS_CHECK"


def _build_workbook_copy(tmp_path: Path, name: str) -> Path:
    workbook_path = tmp_path / name
    shutil.copy(TEMPLATE_WORKBOOK_PATH, workbook_path)
    return workbook_path


def _validate(path: Path):
    workbook = load_workbook(path, data_only=False)
    intro = workbook["INTRO"]
    country = str(intro["E33"].value)
    year = int(float(str(intro["E35"].value)))
    return jrsm.validate_jrsm_workbook(
        workbook_path=path,
        country=country,
        year_for_request_of_medicine=year,
    )


def test_missing_reference_doc_emits_spec_not_configured(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "missing_ref.xlsx")
    missing_reference = tmp_path / "missing_reference_doc.md"
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", missing_reference)

    result = _validate(workbook_path)
    assert any(finding.rule_id == STORY10_RULE_SPEC and finding.severity == "info" for finding in result.findings)


def test_missing_summary_or_shipment_sheet_emits_basic_access_warn(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "missing_sheet.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# placeholder\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)

    workbook = load_workbook(workbook_path, data_only=False)
    summary_sheet = workbook["SUMMARY"]
    workbook.remove(summary_sheet)
    workbook.save(workbook_path)

    result = _validate(workbook_path)
    access_findings = [finding for finding in result.findings if finding.rule_id == STORY10_RULE_ACCESS]
    assert any(finding.sheet == "SUMMARY" and finding.severity == "warn" for finding in access_findings)
    assert any(finding.sheet == "SHIPMENT" and finding.severity == "info" for finding in access_findings)


def test_anchor_cell_not_addressable_emits_basic_access_warn(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "missing_anchor.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# placeholder\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)

    original_is_anchor_addressable = jrsm._is_anchor_addressable

    def fake_is_anchor_addressable(sheet, cell):
        if sheet.title == "SUMMARY" and cell == "G12":
            return False
        return original_is_anchor_addressable(sheet, cell)

    monkeypatch.setattr(jrsm, "_is_anchor_addressable", fake_is_anchor_addressable)

    result = _validate(workbook_path)
    assert any(
        finding.rule_id == STORY10_RULE_ACCESS
        and finding.sheet == "SUMMARY"
        and finding.severity == "warn"
        and "SUMMARY!G12" in (finding.actual or "")
        for finding in result.findings
    )


def test_normal_workbook_emits_basic_access_pass_findings(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "normal.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# placeholder\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)

    result = _validate(workbook_path)
    access_findings = [finding for finding in result.findings if finding.rule_id == STORY10_RULE_ACCESS]
    summary_pass = [finding for finding in access_findings if finding.sheet == "SUMMARY" and finding.severity == "info"]
    shipment_pass = [finding for finding in access_findings if finding.sheet == "SHIPMENT" and finding.severity == "info"]
    assert summary_pass
    assert shipment_pass
