"""Story 10b tests for detailed SUMMARY/SHIPMENT validation."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from espen_sql_api.jap_validation.validators import jrsm

TEMPLATE_WORKBOOK_PATH = Path("docs/JRSM-template-all-sheets.xlsx")
RULE_REFERENCE_APPLIED = "JRSM.SUMMARY_SHIPMENT.REFERENCE_APPLIED"
RULE_SUMMARY_FORMULA_REQUIRED = "JRSM.SUMMARY.FORMULA_REQUIRED"
RULE_SUMMARY_FORMULA_MATCH = "JRSM.SUMMARY.FORMULA_MATCH"
RULE_TABLE_BLOCK_REQUIRED = "JRSM.SUMMARY_SHIPMENT.TABLE_BLOCK_REQUIRED"


def _build_workbook_copy(tmp_path: Path, name: str) -> Path:
    workbook_path = tmp_path / name
    shutil.copy(TEMPLATE_WORKBOOK_PATH, workbook_path)
    return workbook_path


def _validate(path: Path):
    workbook = load_workbook(path, data_only=False)
    intro = workbook["INTRO"]
    country = str(intro["E33"].value)
    year = int(float(str(intro["E35"].value)))
    return jrsm.validate_jrsm_workbook(
        workbook_path=path,
        country=country,
        year_for_request_of_medicine=year,
    )


def _mutate_cell(path: Path, sheet_name: str, cell_address: str, value: object) -> None:
    workbook = load_workbook(path, data_only=False)
    workbook[sheet_name][cell_address] = value
    workbook.save(path)


def test_story10b_reference_applied_info_finding(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "story10b_reference.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# detailed reference\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)

    result = _validate(workbook_path)
    assert any(
        finding.rule_id == RULE_REFERENCE_APPLIED
        and finding.severity == "info"
        for finding in result.findings
    )


def test_summary_non_overlap_formula_required_warn(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "story10b_summary_required.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# detailed reference\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)
    _mutate_cell(workbook_path, "SUMMARY", "C6", "")

    result = _validate(workbook_path)
    assert any(
        finding.rule_id == RULE_SUMMARY_FORMULA_REQUIRED
        and finding.sheet == "SUMMARY"
        and finding.cell == "C6"
        and finding.severity == "warn"
        for finding in result.findings
    )


def test_story9_overlap_formula_mismatch_is_deduped(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "story10b_dedup.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# detailed reference\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)
    _mutate_cell(workbook_path, "SUMMARY", "G12", "=0")

    result = _validate(workbook_path)
    assert any(
        finding.rule_id == "JRSM.FORMULA.GOVERNED_FORMULA_MATCH"
        and finding.sheet == "SUMMARY"
        and finding.cell == "G12"
        for finding in result.findings
    )
    assert not any(
        finding.rule_id == RULE_SUMMARY_FORMULA_MATCH
        and finding.sheet == "SUMMARY"
        and finding.cell == "G12"
        for finding in result.findings
    )


def test_summary_intervention_block_missing_emits_warn(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "story10b_summary_block.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# detailed reference\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)

    # Trigger intervention-date block for row 43 by forcing requested value > 0.
    _mutate_cell(workbook_path, "SUMMARY", "G12", 1)

    result = _validate(workbook_path)
    assert any(
        finding.rule_id == RULE_TABLE_BLOCK_REQUIRED
        and finding.sheet == "SUMMARY"
        and finding.severity == "warn"
        and "SUMMARY.BLOCK.INTERVENTION_DATES.R43" in finding.message
        for finding in result.findings
    )


def test_shipment_consignee_block_missing_emits_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "story10b_shipment_block.xlsx")
    configured_reference = tmp_path / "configured_reference.md"
    configured_reference.write_text("# detailed reference\n", encoding="utf-8")
    monkeypatch.setattr(jrsm, "SUMMARY_SHIPMENT_REFERENCE_DOC_PATH", configured_reference)

    # Activate consignee-set trigger while leaving required recipient fields blank.
    _mutate_cell(workbook_path, "SHIPMENT", "B8", "ALB for LF")

    result = _validate(workbook_path)
    assert any(
        finding.rule_id == RULE_TABLE_BLOCK_REQUIRED
        and finding.sheet == "SHIPMENT"
        and finding.severity == "error"
        and "SHIPMENT.BLOCK.CONSIGNEE_SET_1" in finding.message
        for finding in result.findings
    )
