"""Story 9 tests for JRSM formula governance submitter-day scope."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from espen_sql_api.jap_validation import specs_loader
from espen_sql_api.jap_validation.validators.jrsm import validate_jrsm_workbook

TEMPLATE_WORKBOOK_PATH = Path("docs/JRSM-template-all-sheets.xlsx")
STORY9_RULE_IDS = {
    "JRSM.FORMULA.GOVERNED_FORMULA_PRESENT",
    "JRSM.FORMULA.GOVERNED_FORMULA_MATCH",
    "JRSM.FORMULA.PROTECTED_CELL_ENFORCEMENT",
    "JRSM.FORMULA.EDITABLE_CELL_OVERRIDE_DETECTED",
    "JRSM.FORMULA.SPEC_NOT_CONFIGURED",
    "JRSM.CELL_COLOR.SPEC_NOT_CONFIGURED",
}


def _build_workbook_copy(tmp_path: Path, name: str = "story9.xlsx") -> Path:
    workbook_path = tmp_path / name
    shutil.copy(TEMPLATE_WORKBOOK_PATH, workbook_path)
    return workbook_path


def _validate(workbook_path: Path):
    intro = load_workbook(workbook_path, data_only=False)["INTRO"]
    country = str(intro["E33"].value)
    year = int(float(str(intro["E35"].value)))
    return validate_jrsm_workbook(
        workbook_path=workbook_path,
        country=country,
        year_for_request_of_medicine=year,
    )


def _story9_findings(result):
    return [finding for finding in result.findings if finding.rule_id in STORY9_RULE_IDS]


def _mutate_cell(path: Path, sheet_name: str, cell: str, value: object) -> None:
    workbook = load_workbook(path, data_only=False)
    workbook[sheet_name][cell] = value
    workbook.save(path)


def test_story9_happy_path_has_no_formula_governance_failures(tmp_path: Path):
    workbook_path = _build_workbook_copy(tmp_path, "happy.xlsx")
    result = _validate(workbook_path)
    story9_findings = _story9_findings(result)
    assert not [finding for finding in story9_findings if finding.severity in {"error", "warn"}]


def test_country_info_formula_governance_literal_and_mismatch(tmp_path: Path):
    protected_path = _build_workbook_copy(tmp_path, "country_info_protected.xlsx")
    _mutate_cell(protected_path, "COUNTRY_INFO", "A10", "literal-country")
    protected_result = _validate(protected_path)
    protected_findings = _story9_findings(protected_result)
    assert any(
        finding.rule_id == "JRSM.FORMULA.PROTECTED_CELL_ENFORCEMENT"
        and finding.sheet == "COUNTRY_INFO"
        and finding.cell == "A10"
        and finding.severity == "error"
        for finding in protected_findings
    )

    editable_path = _build_workbook_copy(tmp_path, "country_info_editable.xlsx")
    _mutate_cell(editable_path, "COUNTRY_INFO", "E10", "literal-editable")
    editable_result = _validate(editable_path)
    editable_findings = _story9_findings(editable_result)
    assert any(
        finding.rule_id == "JRSM.FORMULA.EDITABLE_CELL_OVERRIDE_DETECTED"
        and finding.sheet == "COUNTRY_INFO"
        and finding.cell == "E10"
        and finding.severity == "warn"
        for finding in editable_findings
    )


def test_alb_mbd_formula_mismatch_is_error(tmp_path: Path):
    workbook_path = _build_workbook_copy(tmp_path, "alb_mbd_mismatch.xlsx")
    _mutate_cell(workbook_path, "ALB_MBD", "A10", '=IF(1=1,"override","override")')
    result = _validate(workbook_path)
    assert any(
        finding.rule_id == "JRSM.FORMULA.GOVERNED_FORMULA_MATCH"
        and finding.sheet == "ALB_MBD"
        and finding.cell == "A10"
        and finding.severity == "error"
        for finding in _story9_findings(result)
    )


def test_pzq_formula_mismatch_is_warn_for_editable_column(tmp_path: Path):
    workbook_path = _build_workbook_copy(tmp_path, "pzq_mismatch.xlsx")
    _mutate_cell(workbook_path, "PZQ", "E10", "=IF(1=1,123,456)")
    result = _validate(workbook_path)
    assert any(
        finding.rule_id == "JRSM.FORMULA.GOVERNED_FORMULA_MATCH"
        and finding.sheet == "PZQ"
        and finding.cell == "E10"
        and finding.severity == "warn"
        for finding in _story9_findings(result)
    )


def test_ivm_summary_and_shipment_governed_cells(tmp_path: Path):
    ivm_path = _build_workbook_copy(tmp_path, "ivm_mismatch.xlsx")
    _mutate_cell(ivm_path, "IVM", "A10", '=IF(1=1,"bad","bad")')
    ivm_result = _validate(ivm_path)
    assert any(
        finding.rule_id == "JRSM.FORMULA.GOVERNED_FORMULA_MATCH"
        and finding.sheet == "IVM"
        and finding.cell == "A10"
        and finding.severity == "error"
        for finding in _story9_findings(ivm_result)
    )

    summary_path = _build_workbook_copy(tmp_path, "summary_mismatch.xlsx")
    _mutate_cell(summary_path, "SUMMARY", "G12", "=0")
    summary_result = _validate(summary_path)
    assert any(
        finding.rule_id == "JRSM.FORMULA.GOVERNED_FORMULA_MATCH"
        and finding.sheet == "SUMMARY"
        and finding.cell == "G12"
        and finding.severity == "error"
        for finding in _story9_findings(summary_result)
    )

    shipment_path = _build_workbook_copy(tmp_path, "shipment_mismatch.xlsx")
    _mutate_cell(shipment_path, "SHIPMENT", "C4", "=0")
    shipment_result = _validate(shipment_path)
    shipment_findings = _story9_findings(shipment_result)
    assert any(
        finding.rule_id == "JRSM.FORMULA.GOVERNED_FORMULA_MATCH"
        and finding.sheet == "SHIPMENT"
        and finding.cell == "C4"
        and finding.severity == "error"
        for finding in shipment_findings
    )
    assert not any(
        finding.sheet == "SHIPMENT" and finding.cell == "D4"
        for finding in shipment_findings
    )


def test_story9_exclusions_do_not_emit_governance_findings(tmp_path: Path):
    workbook_path = _build_workbook_copy(tmp_path, "exclusions.xlsx")
    workbook = load_workbook(workbook_path, data_only=False)
    workbook["DEC"]["A10"] = "literal-dec"
    workbook["IVM+"]["A10"] = "literal-ivmplus"
    workbook["DATA_POLICY"]["A1"] = "literal-data-policy"
    workbook["SUMMARY"]["D12"] = "literal-summary-excluded"
    workbook.save(workbook_path)

    result = _validate(workbook_path)
    story9_findings = _story9_findings(result)
    assert not any(finding.sheet in {"DEC", "IVM+", "DATA_POLICY"} for finding in story9_findings)
    assert not any(finding.sheet == "SUMMARY" and finding.cell == "D12" for finding in story9_findings)
    assert not any(finding.sheet == "SHIPMENT" and finding.cell == "D4" for finding in story9_findings)


def test_story9_missing_specs_emit_info_fallback_findings(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    workbook_path = _build_workbook_copy(tmp_path, "missing_specs.xlsx")
    missing_formula_spec = tmp_path / "missing_formula_spec.json"
    missing_color_spec = tmp_path / "missing_color_spec.yaml"
    monkeypatch.setattr(specs_loader, "JRSM_FORMULA_SPEC_PATH", missing_formula_spec)
    monkeypatch.setattr(specs_loader, "JRSM_CELL_COLOR_MAP_PATH", missing_color_spec)

    result = _validate(workbook_path)
    story9_findings = _story9_findings(result)
    assert any(finding.rule_id == "JRSM.FORMULA.SPEC_NOT_CONFIGURED" for finding in story9_findings)
    assert any(finding.rule_id == "JRSM.CELL_COLOR.SPEC_NOT_CONFIGURED" for finding in story9_findings)
