"""JRSM deterministic validator for Story 4 workbook/template checks."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import re
from typing import Any
import unicodedata

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, coordinate_to_tuple, get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..models import Finding, ValidationEngineResult
from ..specs_loader import load_jrsm_cell_color_map, load_jrsm_formula_spec

REQUIRED_SHEETS = [
    "INTRO",
    "COUNTRY_INFO",
    "DEC",
    "ALB_MBD",
    "PZQ",
    "IVM",
    "IVM+",
    "SUMMARY",
    "SHIPMENT",
    "DATA_POLICY",
]
LANGUAGE_PROFILE_RULE_ID = "JRSM.TEMPLATE.LANGUAGE_PROFILE"
VERSION_MARKER_RULE_ID = "JRSM.TEMPLATE.VERSION_MARKER"
SUPPORTED_TEMPLATE_STATUS = "supported"
UNSUPPORTED_VERSION_STATUS = "unsupported_version"
UNDETECTED_LANGUAGE_STATUS = "undetected_language"
JRSM_LANGUAGE_PROFILES = {
    "en": {
        "label": "English",
        "supported_markers": ("Joint request for selected PC medicines v.4.4",),
        "marker_prefixes": ("Joint request for selected PC medicines",),
        "fallback_header_tokens": {"H7": "LF", "I7": "Oncho", "J7": "STH", "K7": "SCH"},
        "formula_literals": {
            "Endemic but PC is not required": "Endemic but PC is not required",
            "Non-endemic": "Non-endemic",
            "Endemic": "Endemic",
            "Unknown": "Unknown",
            "Stopped": "Stopped",
            "Surveillance": "Surveillance",
            "Not required": "Not required",
            "Treat with DEC": "Treat with DEC",
            "Treat with IVM": "Treat with IVM",
            "Yes": "Yes",
        },
        "semantic_literals": {
            "endemic": {"Endemic"},
            "non_endemic": {"Non-endemic"},
            "endemic_but_pc_not_required": {"Endemic but PC is not required"},
        },
    },
    "fr": {
        "label": "French",
        "supported_markers": ("Formulaire de demande commune de médicaments pour CP v.4.5",),
        "marker_prefixes": ("Formulaire de demande commune de médicaments pour CP",),
        "fallback_header_tokens": {"H7": "FL", "I7": "Oncho", "J7": "STH", "K7": "SCH"},
        "formula_literals": {
            "Endemic but PC is not required": "Endémique mais pas de CP",
            "Non-endemic": "Non endémique",
            "Endemic": "Endémique",
            "Unknown": "Inconnu",
            "Stopped": "Interrompu",
            "Surveillance": "Surveillance",
            "Not required": "Non requis",
            "Treat with DEC": "Traiter avec DEC",
            "Treat with IVM": "Traiter avec IVM",
            "Yes": "Oui",
        },
        "semantic_literals": {
            "endemic": {"Endémique"},
            "non_endemic": {"Non endémique"},
            "endemic_but_pc_not_required": {"Endémique mais pas de CP"},
        },
    },
    "es": {
        "label": "Spanish",
        "supported_markers": ("Solicitud conjunta de medicamentos seleccionados para QP v.4.5",),
        "marker_prefixes": ("Solicitud conjunta de medicamentos seleccionados para QP",),
        "fallback_header_tokens": {"H7": "FL", "I7": "Onco", "J7": "HTS", "K7": "ESQ"},
        "formula_literals": {
            "Endemic but PC is not required": "Endémico pero no requiere QP",
            "Non-endemic": "No endémico",
            "Endemic": "Endémico",
            "Unknown": "Desconocido",
            "Stopped": "Suspendida",
            "Surveillance": "Vigilancia",
            "Not required": "No requerido",
            "Treat with DEC": "Tratar con DEC",
            "Treat with IVM": "Tratar con IVM",
            "Yes": "Sí",
        },
        "semantic_literals": {
            "endemic": {"Endémico"},
            "non_endemic": {"No endémico"},
            "endemic_but_pc_not_required": {"Endémico pero no requiere QP"},
        },
    },
}
EXPECTED_VERSION_MARKER = JRSM_LANGUAGE_PROFILES["en"]["marker_prefixes"][0]
MODULE_VISIBILITY_SHEETS = ["DEC", "IVM", "IVM+", "ALB_MBD", "PZQ"]
CORE_EXPECTED_VISIBLE_SHEETS = ["INTRO", "COUNTRY_INFO", "SUMMARY", "SHIPMENT"]
REQUIRED_INTRO_FIELD_RANGES = [("E37", "E45"), ("E48", "E51")]
REQUIRED_INTRO_CELLS = [
    "B2",
    "E33",
    "E35",
    "E37",
    "E38",
    "E39",
    "E40",
    "E41",
    "E42",
    "E43",
    "E44",
    "E45",
    "E48",
    "E49",
    "E50",
    "E51",
]


def _expand_cell_range(start: str, end: str) -> list[str]:
    start_row, start_col = coordinate_to_tuple(start)
    end_row, end_col = coordinate_to_tuple(end)
    if start_col != end_col:
        raise ValueError(f"Only single-column ranges are supported: {start}:{end}")

    column_letter = start[0]
    return [f"{column_letter}{row}" for row in range(start_row, end_row + 1)]


REQUIRED_INTRO_FIELD_CELLS = [
    cell
    for range_start, range_end in REQUIRED_INTRO_FIELD_RANGES
    for cell in _expand_cell_range(range_start, range_end)
]
REQUIRED_INTRO_VALUE_CELLS = ["E37", "E39", "E41", "E43", "E45", "E48", "E49", "E50", "E51"]
MODULE_COMPLETENESS_SPECS = [
    (
        "ALB_MBD",
        ["G", "J", "Q", "T"],
        "JRSM.ALB_MBD.REQUIRED_COLUMNS_G_J_Q_T",
        "JRSM.ALB_MBD.NOT_APPLICABLE",
    ),
    (
        "PZQ",
        ["I", "J", "K", "L", "Q"],
        "JRSM.PZQ.REQUIRED_COLUMNS_I_TO_L_AND_Q",
        "JRSM.PZQ.NOT_APPLICABLE",
    ),
    (
        "IVM",
        ["M"],
        "JRSM.IVM.REQUIRED_COLUMN_M",
        "JRSM.IVM.NOT_APPLICABLE",
    ),
]
CONTEXT_OPTIONAL_COUNTRY_INFO_OUTPUT_COLUMNS = {"AA", "AB", "AC", "AD"}
COLUMN_LABEL_OVERRIDES = {
    "ALB_MBD": {
        "G": "Target population for STH / PreSAC / Rounds",
        "J": "WRA / Rounds",
        "Q": "Remaining in stock (STH)",
        "T": "Remaining in stock",
    },
    "PZQ": {
        "I": "Hot spots / Planned",
        "J": "Population",
        "K": "Previous year MDA / Rounds",
        "L": "Age groups",
        "Q": "Remaining in stock",
    },
    "IVM": {
        "M": "Remaining in stock",
    },
}


def _summary_shipment_reference_doc_path_for_module_path(module_path: Path) -> Path:
    jap_validation_dir = module_path.parent.parent
    return jap_validation_dir / "specs/jrsm_summary_shipment_validation_reference.md"


SUMMARY_SHIPMENT_REFERENCE_DOC_PATH = (
    _summary_shipment_reference_doc_path_for_module_path(Path(__file__).resolve(strict=True))
)
SUMMARY_SHIPMENT_ANCHOR_CELLS = {
    "SUMMARY": ["A1", "G12", "B61", "C61"],
    "SHIPMENT": ["A1", "C4", "H4"],
}
STORY10B_SUMMARY_CANONICAL_FORMULAS = {
    "C6": '=IF(INTRO!$E$33<>0,INTRO!$E$33," ")',
    "H6": '=IF(INTRO!$E$35<>0,INTRO!$E$35," ")',
    "G12": "=IF(D12>(E12+F12),D12-E12-F12, 0)",
    "G13": "=IF(D13>(E13+F13),D13-E13-F13, 0)",
    "G14": "=IF(D14>(E14+F14),D14-E14-F14, 0)",
    "G16": "=IF(D16>(E16+F16),D16-E16-F16, 0)",
    "G17": "=IF(D17>(E17+F17),D17-E17-F17, 0)",
    "G18": "=IF(D18>(E18+F18),D18-E18-F18, 0)",
    "G19": "=IF(D19>(E19+F19),D19-E19-F19, 0)",
    "G22": "=IF(D22>(E22+F22),D22-E22-F22, 0)",
    "G23": "=IF(D23>(E23+F23),D23-E23-F23, 0)",
    "G24": "=IF(D24>(E24+F24),D24-E24-F24, 0)",
    "G26": "=IF(D26>(E26+F26),D26-E26-F26, 0)",
    "G27": "=IF(D27>(E27+F27),D27-E27-F27, 0)",
    "G28": "=IF(D28>(E28+F28),D28-E28-F28, 0)",
    "B61": "=B33",
    "C61": "=SUM(F33:H33)",
    "B62": "=B34",
    "C62": "=SUM(F34:H34)",
    "B63": "=B35",
    "C63": "=SUM(F35:H35)",
    "B64": "=B36",
    "C64": "=SUM(F36:H36)",
}
STORY10B_SHIPMENT_CANONICAL_FORMULAS = {
    "C4": '=IF(INTRO!$E$33<>0,INTRO!$E$33," ")',
    "H4": '=IF(INTRO!$E$35<>0,INTRO!$E$35," ")',
}
STORY10B_FORMULA_NON_OVERLAP_WARN_CELLS = {"SUMMARY": {"C6", "H6"}}
STORY10B_STORY9_FORMULA_OVERLAP = {
    "SUMMARY": {
        "G12",
        "G13",
        "G14",
        "G16",
        "G17",
        "G18",
        "G19",
        "G22",
        "G23",
        "G24",
        "G26",
        "G27",
        "G28",
        "B61",
        "C61",
        "B62",
        "C62",
        "B63",
        "C63",
        "B64",
        "C64",
    },
    "SHIPMENT": {"C4", "H4"},
}
STORY10B_FORMULA_REQUIRED_RULES = {
    "JRSM.FORMULA.GOVERNED_FORMULA_PRESENT",
    "JRSM.FORMULA.PROTECTED_CELL_ENFORCEMENT",
    "JRSM.FORMULA.EDITABLE_CELL_OVERRIDE_DETECTED",
}


def _has_cell_address(sheet: Worksheet, cell_address: str) -> bool:
    try:
        _ = sheet[cell_address]
    except Exception:
        return False
    return True


def _build_finding(
    findings: list[Finding],
    *,
    rule_id: str,
    severity: str,
    message: str,
    recommendation: str,
    sheet: str | None = None,
    cell: str | None = None,
    expected: str | None = None,
    actual: str | None = None,
) -> None:
    issue_id = f"F-{len(findings) + 1:04d}"
    findings.append(
        Finding(
            issue_id=issue_id,
            rule_id=rule_id,
            severity=severity,  # type: ignore[arg-type]
            sheet=sheet,
            cell=cell,
            expected=expected,
            actual=actual,
            message=message,
            recommendation=recommendation,
        )
    )


def _derive_outcome(findings: list[Finding]) -> str:
    return "fail" if any(finding.severity == "error" for finding in findings) else "pass"


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split()).casefold()


def _canonicalize_text(value: Any) -> str:
    if value is None:
        return ""
    normalized = unicodedata.normalize("NFKD", str(value))
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    return " ".join(ascii_text.split()).casefold()


def _normalize_year(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    if isinstance(value, (date, datetime)):
        return value.year

    text_value = " ".join(str(value).split()).replace(",", "")
    if not text_value:
        return None
    try:
        numeric_value = float(text_value)
    except ValueError:
        return None
    if not numeric_value.is_integer():
        return None
    return int(numeric_value)


def _normalize_positive_int(value: Any) -> int | None:
    normalized = _normalize_year(value)
    if normalized is None:
        return None
    return normalized if normalized > 0 else None


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _header_text(sheet: Worksheet, column_letter: str) -> str:
    values: list[str] = []
    for row in (6, 7, 8):
        cell_address = f"{column_letter}{row}"
        value = sheet[cell_address].value
        if _is_blank(value):
            for merged_range in sheet.merged_cells.ranges:
                if cell_address not in merged_range:
                    continue
                value = sheet[merged_range.start_cell.coordinate].value
                break
        if _is_blank(value):
            continue
        text_value = str(value).strip()
        if text_value in values:
            continue
        values.append(text_value)
    return " / ".join(values)


def _column_label(sheet: Worksheet, column_letter: str) -> str:
    override = COLUMN_LABEL_OVERRIDES.get(sheet.title, {}).get(column_letter)
    if override:
        return override
    return _header_text(sheet, column_letter) or column_letter


def _format_column_descriptor(sheet: Worksheet, columns: list[str]) -> str:
    if len(columns) == 1:
        column = columns[0]
        return f"column {column} ({_column_label(sheet, column)})"

    labels = ", ".join(_column_label(sheet, column) for column in columns)
    return f"columns {columns[0]}-{columns[-1]} ({labels})"


def _format_grouped_cell_range(columns: list[str], start_row: int, end_row: int) -> str:
    if len(columns) == 1:
        return f"{columns[0]}{start_row}:{columns[0]}{end_row}"
    return f"{columns[0]}{start_row}:{columns[-1]}{end_row}"


def _group_contiguous_columns(columns: list[str]) -> list[list[str]]:
    if not columns:
        return []

    groups: list[list[str]] = [[columns[0]]]
    for column in columns[1:]:
        previous = groups[-1][-1]
        if column_index_from_string(column) == column_index_from_string(previous) + 1:
            groups[-1].append(column)
        else:
            groups.append([column])
    return groups


def _is_output_column_expected(
    *,
    header_text: str,
    has_lf: bool,
    has_oncho: bool,
    has_sth: bool,
    has_sch: bool,
    has_ida: bool | None,
) -> bool:
    if not header_text:
        return False

    header_lower = header_text.casefold()
    keyword_map = {
        "lf": has_lf,
        "lymphatic": has_lf,
        "filaria": has_lf,
        "oncho": has_oncho,
        "onchocer": has_oncho,
        "sth": has_sth,
        "soil-transmitted": has_sth,
        "soil transmitted": has_sth,
        "sch": has_sch,
        "schisto": has_sch,
        "ida": bool(has_ida),
    }

    matched_any = False
    matched_expected = False
    for keyword, enabled in keyword_map.items():
        if keyword in header_lower:
            matched_any = True
            matched_expected = matched_expected or enabled
    if matched_any:
        return matched_expected
    return True


def _is_column_hidden(sheet: Worksheet, column_letter: str) -> bool:
    column_dimension = sheet.column_dimensions.get(column_letter)
    return bool(column_dimension and column_dimension.hidden)


def _cell_column_letter(cell_address: str) -> str:
    _, column_index = coordinate_to_tuple(cell_address)
    return get_column_letter(column_index)


def _has_any_value_in_columns(*, sheet: Worksheet, columns: set[str], rows: list[int]) -> bool:
    return any(not _is_blank(sheet[f"{column}{row}"].value) for column in columns for row in rows)


def _is_context_optional_country_info_output_column_required(
    *,
    sheet: Worksheet,
    column_letter: str,
    active_rows: list[int],
) -> bool:
    if column_letter not in CONTEXT_OPTIONAL_COUNTRY_INFO_OUTPUT_COLUMNS:
        return True
    return _has_any_value_in_columns(
        sheet=sheet,
        columns=CONTEXT_OPTIONAL_COUNTRY_INFO_OUTPUT_COLUMNS,
        rows=active_rows,
    )


def _build_reviewable_sheet_visibility(
    *,
    workbook: Workbook,
    expected_sheet_visibility: dict[str, bool],
) -> dict[str, bool]:
    reviewable_sheet_visibility: dict[str, bool] = {}
    for sheet_name in REQUIRED_SHEETS:
        if sheet_name not in workbook.sheetnames:
            reviewable_sheet_visibility[sheet_name] = False
            continue

        if sheet_name in CORE_EXPECTED_VISIBLE_SHEETS:
            reviewable_sheet_visibility[sheet_name] = True
            continue

        if sheet_name in MODULE_VISIBILITY_SHEETS:
            reviewable_sheet_visibility[sheet_name] = (
                bool(expected_sheet_visibility.get(sheet_name, False))
                and workbook[sheet_name].sheet_state == "visible"
            )
            continue

        reviewable_sheet_visibility[sheet_name] = workbook[sheet_name].sheet_state == "visible"
    return reviewable_sheet_visibility


def _is_formula_value(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _normalize_formula(formula: str) -> str:
    # Normalize for deterministic comparison across spacing/case/$-anchor variants.
    normalized = re.sub(r"\s+", "", formula).upper()
    normalized = normalized.replace("$", "")
    # Strip external workbook scoping tokens observed in submitted files, e.g. [1]INTRO!...
    return re.sub(r"\[\d+\](?=[A-Z_][A-Z0-9_]*!)", "", normalized)


def _to_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _is_yes(value: Any) -> bool:
    return _canonicalize_text(value) in {"yes", "oui", "si"}


def _true_if_matches_any(value: Any, candidates: set[str]) -> bool:
    normalized_value = _canonicalize_text(value)
    return bool(normalized_value) and normalized_value in {_canonicalize_text(candidate) for candidate in candidates}


def _template_profile_supported_markers(profile_code: str) -> set[str]:
    return {
        _canonicalize_text(marker)
        for marker in JRSM_LANGUAGE_PROFILES[profile_code]["supported_markers"]
    }


def _template_profile_marker_prefixes(profile_code: str) -> set[str]:
    return {
        _canonicalize_text(marker_prefix)
        for marker_prefix in JRSM_LANGUAGE_PROFILES[profile_code]["marker_prefixes"]
    }


def _detect_language_from_headers(workbook: Workbook) -> str | None:
    if "COUNTRY_INFO" not in workbook.sheetnames:
        return None

    country_info = workbook["COUNTRY_INFO"]
    for profile_code, profile in JRSM_LANGUAGE_PROFILES.items():
        expected_tokens = profile.get("fallback_header_tokens", {})
        if not expected_tokens:
            continue
        if all(
            _canonicalize_text(country_info[cell_address].value) == _canonicalize_text(expected_value)
            for cell_address, expected_value in expected_tokens.items()
        ):
            return profile_code
    return None


def _detect_template_profile(workbook: Workbook, intro_sheet: Worksheet | None) -> dict[str, str | None]:
    marker_text = ""
    if intro_sheet is not None and _is_anchor_addressable(intro_sheet, "B2"):
        b2_value = intro_sheet["B2"].value
        marker_text = str(b2_value).strip() if b2_value is not None else ""
    marker_normalized = _canonicalize_text(marker_text)

    if marker_normalized:
        for profile_code in JRSM_LANGUAGE_PROFILES:
            if marker_normalized in _template_profile_supported_markers(profile_code):
                return {
                    "status": SUPPORTED_TEMPLATE_STATUS,
                    "language": profile_code,
                    "marker_text": marker_text,
                }
        for profile_code in JRSM_LANGUAGE_PROFILES:
            if any(
                marker_normalized.startswith(marker_prefix)
                for marker_prefix in _template_profile_marker_prefixes(profile_code)
            ):
                return {
                    "status": UNSUPPORTED_VERSION_STATUS,
                    "language": profile_code,
                    "marker_text": marker_text,
                }

    guessed_language = _detect_language_from_headers(workbook)
    return {
        "status": UNDETECTED_LANGUAGE_STATUS,
        "language": guessed_language,
        "marker_text": marker_text or None,
    }


def _emit_template_profile_finding(findings: list[Finding], profile_state: dict[str, str | None]) -> None:
    status = profile_state["status"]
    language = profile_state.get("language")
    marker_text = profile_state.get("marker_text")
    language_label = JRSM_LANGUAGE_PROFILES.get(language or "", {}).get("label")

    if status == SUPPORTED_TEMPLATE_STATUS and language and language_label:
        _build_finding(
            findings,
            rule_id=LANGUAGE_PROFILE_RULE_ID,
            severity="info",
            message=f"Detected supported {language_label} JRSM template profile.",
            recommendation="No action required.",
            sheet="INTRO",
            cell="B2",
            expected="Supported JRSM language profile",
            actual=language,
        )
        return

    if status == UNSUPPORTED_VERSION_STATUS and language and language_label:
        _build_finding(
            findings,
            rule_id=LANGUAGE_PROFILE_RULE_ID,
            severity="warn",
            message=f"Detected {language_label} JRSM workbook language, but the template version is outside the supported baseline.",
            recommendation=f"Use the current supported {language_label} JRSM template before relying on full deterministic validation.",
            sheet="INTRO",
            cell="B2",
            expected=JRSM_LANGUAGE_PROFILES[language]["supported_markers"][0],
            actual=marker_text or "BLANK_OR_MISSING",
        )
        return

    _build_finding(
        findings,
        rule_id=LANGUAGE_PROFILE_RULE_ID,
        severity="warn",
        message="JRSM workbook language/profile could not be confidently detected.",
        recommendation="Use a current supported English, French, or Spanish JRSM template with the standard INTRO!B2 marker.",
        sheet="INTRO",
        cell="B2",
        expected="Supported JRSM template marker in INTRO!B2",
        actual=marker_text or "BLANK_OR_MISSING",
    )


def _formula_literal_map_for_language(template_language: str | None) -> dict[str, str]:
    if template_language is None:
        return {}
    profile = JRSM_LANGUAGE_PROFILES.get(template_language)
    if not profile:
        return {}
    return profile["formula_literals"]


def _localize_formula_literals(formula: str, *, template_language: str | None) -> str:
    literal_map = _formula_literal_map_for_language(template_language)
    if not literal_map:
        return formula

    localized_formula = formula
    for english_literal in sorted(literal_map, key=len, reverse=True):
        localized_formula = localized_formula.replace(
            f'"{english_literal}"',
            f'"{literal_map[english_literal]}"',
        )
    return localized_formula


def _matches_template_semantic_literal(
    value: Any,
    *,
    template_language: str | None,
    semantic_key: str,
) -> bool:
    if template_language is None or template_language not in JRSM_LANGUAGE_PROFILES:
        return False
    candidates = JRSM_LANGUAGE_PROFILES[template_language]["semantic_literals"].get(semantic_key, set())
    return _true_if_matches_any(value, candidates)


def _translate_formula_to_target(*, formula: str, source_cell: str, target_cell: str) -> str:
    if source_cell == target_cell:
        return formula
    try:
        return Translator(formula, origin=source_cell).translate_formula(target_cell)
    except Exception:
        return formula


def _format_actual_value(value: Any) -> str:
    if _is_blank(value):
        return "BLANK_OR_MISSING"
    return str(value)


def _governance_severity(governance_class: str) -> str:
    return "error" if governance_class == "must_not_change" else "warn"


def _literal_override_rule(governance_class: str) -> str:
    if governance_class == "must_not_change":
        return "JRSM.FORMULA.PROTECTED_CELL_ENFORCEMENT"
    return "JRSM.FORMULA.EDITABLE_CELL_OVERRIDE_DETECTED"


def _governance_recommendation(governance_class: str) -> str:
    if governance_class == "must_not_change":
        return "Restore the canonical template formula for this protected governed cell."
    return "Review editable formula override and restore canonical formula if override is unintended."


def _is_anchor_addressable(sheet: Worksheet, cell_address: str) -> bool:
    try:
        coordinate_to_tuple(cell_address)
        _ = sheet[cell_address]
    except Exception:
        return False
    return True


def _evaluate_story9_formula_governance(
    *,
    workbook: Workbook,
    findings: list[Finding],
    iu_start_row: int,
    iu_end_row: int | None,
    template_language: str | None,
    reviewable_sheet_visibility: dict[str, bool],
) -> None:
    formula_spec, formula_spec_error = load_jrsm_formula_spec()
    cell_color_map, cell_color_map_error = load_jrsm_cell_color_map()
    if formula_spec is None:
        _build_finding(
            findings,
            rule_id="JRSM.FORMULA.SPEC_NOT_CONFIGURED",
            severity="info",
            message="JRSM formula governance spec is not configured; Story 9 enforcement skipped.",
            recommendation="Provide src/espen_sql_api/jap_validation/specs/jrsm_formula_spec.json for Story 9 enforcement.",
            expected="Loadable JSON formula governance spec",
            actual=formula_spec_error or "UNKNOWN_ERROR",
        )

    if cell_color_map is None:
        _build_finding(
            findings,
            rule_id="JRSM.CELL_COLOR.SPEC_NOT_CONFIGURED",
            severity="info",
            message="JRSM cell-color map is not configured; continuing with governance-class defaults.",
            recommendation="Provide src/espen_sql_api/jap_validation/specs/jrsm_cell_color_map.yaml.",
            expected="Loadable YAML cell-color map",
            actual=cell_color_map_error or "UNKNOWN_ERROR",
        )
    else:
        _ = cell_color_map

    if formula_spec is None:
        return

    def evaluate_cell(
        *,
        sheet_name: str,
        cell_address: str,
        governance_class: str,
        expected_formula: str,
    ) -> None:
        if sheet_name not in workbook.sheetnames:
            return

        sheet = workbook[sheet_name]
        if not reviewable_sheet_visibility.get(sheet_name, True):
            return
        if _is_column_hidden(sheet, _cell_column_letter(cell_address)):
            return

        actual_value = sheet[cell_address].value
        severity = _governance_severity(governance_class)
        recommendation = _governance_recommendation(governance_class)

        if not _is_formula_value(actual_value):
            actual_text = _format_actual_value(actual_value)
            localized_expected_formula = _localize_formula_literals(
                expected_formula,
                template_language=template_language,
            )
            _build_finding(
                findings,
                rule_id="JRSM.FORMULA.GOVERNED_FORMULA_PRESENT",
                severity=severity,
                message="Governed cell is missing a formula.",
                recommendation=recommendation,
                sheet=sheet_name,
                cell=cell_address,
                expected="Formula present in governed cell",
                actual=actual_text,
            )
            _build_finding(
                findings,
                rule_id=_literal_override_rule(governance_class),
                severity=severity,
                message="Literal value detected in formula-governed cell.",
                recommendation=recommendation,
                sheet=sheet_name,
                cell=cell_address,
                expected=localized_expected_formula,
                actual=actual_text,
            )
            return

        actual_formula = str(actual_value)
        localized_expected_formula = _localize_formula_literals(
            expected_formula,
            template_language=template_language,
        )
        if _normalize_formula(actual_formula) == _normalize_formula(localized_expected_formula):
            return

        _build_finding(
            findings,
            rule_id="JRSM.FORMULA.GOVERNED_FORMULA_MATCH",
            severity=severity,
            message="Governed cell formula differs from canonical template formula.",
            recommendation=recommendation,
            sheet=sheet_name,
            cell=cell_address,
            expected=localized_expected_formula,
            actual=actual_formula,
        )

    iu_sheet_order = formula_spec.get("iu_sheet_order", [])
    if isinstance(iu_sheet_order, list) and iu_end_row is not None:
        iu_sheets = formula_spec.get("iu_sheets", {})
        if isinstance(iu_sheets, dict):
            for sheet_name in iu_sheet_order:
                if not isinstance(sheet_name, str):
                    continue
                if not reviewable_sheet_visibility.get(sheet_name, True):
                    continue
                sheet_spec = iu_sheets.get(sheet_name)
                if not isinstance(sheet_spec, dict):
                    continue

                canonical_anchor_row = sheet_spec.get("canonical_anchor_row", iu_start_row)
                if not isinstance(canonical_anchor_row, int):
                    canonical_anchor_row = iu_start_row
                canonical_formulas_by_column = sheet_spec.get("canonical_formulas_by_column", {})
                if not isinstance(canonical_formulas_by_column, dict):
                    canonical_formulas_by_column = {}

                for governance_class, columns_key in [
                    ("must_not_change", "must_not_change_columns"),
                    ("editable_override_warn", "editable_override_warn_columns"),
                ]:
                    columns = sheet_spec.get(columns_key, [])
                    if not isinstance(columns, list):
                        continue

                    for column in columns:
                        if not isinstance(column, str):
                            continue
                        canonical_formula = canonical_formulas_by_column.get(column)
                        if not isinstance(canonical_formula, str) or not canonical_formula.startswith("="):
                            continue

                        source_cell = f"{column}{canonical_anchor_row}"
                        for row in range(iu_start_row, iu_end_row + 1):
                            target_cell = f"{column}{row}"
                            expected_formula = _translate_formula_to_target(
                                formula=canonical_formula,
                                source_cell=source_cell,
                                target_cell=target_cell,
                            )
                            evaluate_cell(
                                sheet_name=sheet_name,
                                cell_address=target_cell,
                                governance_class=governance_class,
                                expected_formula=expected_formula,
                            )

    non_iu_sheet_order = formula_spec.get("non_iu_sheet_order", [])
    if isinstance(non_iu_sheet_order, list):
        non_iu_sheets = formula_spec.get("non_iu_sheets", {})
        if isinstance(non_iu_sheets, dict):
            for sheet_name in non_iu_sheet_order:
                if not isinstance(sheet_name, str):
                    continue
                sheet_spec = non_iu_sheets.get(sheet_name)
                if not isinstance(sheet_spec, dict):
                    continue
                canonical_formulas_by_cell = sheet_spec.get("canonical_formulas_by_cell", {})
                if not isinstance(canonical_formulas_by_cell, dict):
                    canonical_formulas_by_cell = {}

                for governance_class, cells_key in [
                    ("must_not_change", "must_not_change_cells"),
                    ("editable_override_warn", "editable_override_warn_cells"),
                ]:
                    cells = sheet_spec.get(cells_key, [])
                    if not isinstance(cells, list):
                        continue

                    for cell_address in cells:
                        if not isinstance(cell_address, str):
                            continue
                        canonical_formula = canonical_formulas_by_cell.get(cell_address)
                        if not isinstance(canonical_formula, str) or not canonical_formula.startswith("="):
                            continue
                        evaluate_cell(
                            sheet_name=sheet_name,
                            cell_address=cell_address,
                            governance_class=governance_class,
                            expected_formula=canonical_formula,
                        )


def _evaluate_story10_summary_shipment_placeholder(
    *,
    workbook: Workbook,
    findings: list[Finding],
) -> None:
    reference_is_configured = SUMMARY_SHIPMENT_REFERENCE_DOC_PATH.exists()
    if not reference_is_configured:
        _build_finding(
            findings,
            rule_id="JRSM.SUMMARY_SHIPMENT.SPEC_NOT_CONFIGURED",
            severity="info",
            message="Detailed SUMMARY/SHIPMENT validation reference is not configured; running placeholder checks only.",
            recommendation=(
                "Add src/espen_sql_api/jap_validation/specs/jrsm_summary_shipment_validation_reference.md "
                "to enable detailed Story 10+ rules."
            ),
            expected=str(SUMMARY_SHIPMENT_REFERENCE_DOC_PATH),
            actual="MISSING_REFERENCE_DOC",
        )

    for sheet_name in ["SUMMARY", "SHIPMENT"]:
        anchor_cells = SUMMARY_SHIPMENT_ANCHOR_CELLS[sheet_name]
        if sheet_name not in workbook.sheetnames:
            _build_finding(
                findings,
                rule_id="JRSM.SUMMARY_SHIPMENT.BASIC_ACCESS_CHECK",
                severity="warn",
                message=f"{sheet_name} basic-access check failed: required sheet is missing.",
                recommendation="Regenerate workbook from standard template and ensure required sheet exists.",
                sheet=sheet_name,
                expected=f"Sheet present with addressable anchors: {', '.join(anchor_cells)}",
                actual="MISSING_SHEET",
            )
            continue

        try:
            sheet = workbook[sheet_name]
        except Exception as exc:
            _build_finding(
                findings,
                rule_id="JRSM.SUMMARY_SHIPMENT.BASIC_ACCESS_CHECK",
                severity="warn",
                message=f"{sheet_name} basic-access check failed: sheet could not be read.",
                recommendation="Verify workbook integrity and resubmit.",
                sheet=sheet_name,
                expected=f"Readable sheet with addressable anchors: {', '.join(anchor_cells)}",
                actual=f"READ_ACCESS_ERROR: {exc}",
            )
            continue

        missing_anchors = [cell for cell in anchor_cells if not _is_anchor_addressable(sheet, cell)]
        if missing_anchors:
            _build_finding(
                findings,
                rule_id="JRSM.SUMMARY_SHIPMENT.BASIC_ACCESS_CHECK",
                severity="warn",
                message=f"{sheet_name} basic-access check failed: anchor cells are not addressable.",
                recommendation="Use a standard JRSM workbook template so placeholder anchor cells are present.",
                sheet=sheet_name,
                expected=f"Addressable anchors: {', '.join(anchor_cells)}",
                actual=", ".join([f"{sheet_name}!{cell}" for cell in missing_anchors]),
            )
            continue

        _build_finding(
            findings,
            rule_id="JRSM.SUMMARY_SHIPMENT.BASIC_ACCESS_CHECK",
            severity="info",
            message=(
                f"{sheet_name} basic-access checks passed."
                if reference_is_configured
                else f"{sheet_name} basic-access placeholder checks passed."
            ),
            recommendation=(
                "No action required."
                if reference_is_configured
                else "No action required until detailed SUMMARY/SHIPMENT rule reference is configured."
            ),
            sheet=sheet_name,
            expected=f"Sheet present and anchors addressable: {', '.join(anchor_cells)}",
            actual="ACCESS_OK",
        )


def _story10b_formula_severity(sheet_name: str, cell_address: str) -> str:
    warn_cells = STORY10B_FORMULA_NON_OVERLAP_WARN_CELLS.get(sheet_name, set())
    return "warn" if cell_address in warn_cells else "error"


def _should_suppress_story10b_formula_finding(
    *,
    findings: list[Finding],
    sheet_name: str,
    cell_address: str,
    check_type: str,
) -> bool:
    if cell_address not in STORY10B_STORY9_FORMULA_OVERLAP.get(sheet_name, set()):
        return False

    expected_story9_rules = (
        STORY10B_FORMULA_REQUIRED_RULES
        if check_type == "required"
        else {"JRSM.FORMULA.GOVERNED_FORMULA_MATCH"}
    )
    return any(
        finding.sheet == sheet_name
        and finding.cell == cell_address
        and finding.rule_id in expected_story9_rules
        for finding in findings
    )


def _emit_story10b_table_block_finding(
    *,
    findings: list[Finding],
    sheet_name: str,
    block_id: str,
    severity: str,
    missing_cells: list[str],
    expected_cells: list[str],
) -> None:
    if not missing_cells:
        return

    _build_finding(
        findings,
        rule_id="JRSM.SUMMARY_SHIPMENT.TABLE_BLOCK_REQUIRED",
        severity=severity,
        message=f"{block_id} required fields are missing or invalid.",
        recommendation="Populate required fields for this SUMMARY/SHIPMENT block and rerun validation.",
        sheet=sheet_name,
        cell=", ".join(missing_cells),
        expected=", ".join([f"{sheet_name}!{cell}" for cell in expected_cells]),
        actual=", ".join([f"{sheet_name}!{cell}" for cell in missing_cells]),
    )


def _evaluate_story10b_summary_table_blocks(*, sheet: Worksheet, findings: list[Finding]) -> None:
    intervention_row_map = [(43, 12), (44, 14), (45, 16), (46, 18), (47, 22), (48, 26)]
    for planning_row, request_row in intervention_row_map:
        requested_value = _to_float(sheet[f"G{request_row}"].value)
        if requested_value is None or requested_value <= 0:
            continue
        expected_cells = [f"C{planning_row}", f"D{planning_row}", f"G{planning_row}", f"H{planning_row}"]
        missing_cells = [cell for cell in expected_cells if _is_blank(sheet[cell].value)]
        _emit_story10b_table_block_finding(
            findings=findings,
            sheet_name="SUMMARY",
            block_id=f"SUMMARY.BLOCK.INTERVENTION_DATES.R{planning_row}",
            severity="warn",
            missing_cells=missing_cells,
            expected_cells=expected_cells,
        )

    for row in [61, 62, 63, 64]:
        treatment_count = _to_float(sheet[f"C{row}"].value)
        if treatment_count is None or treatment_count <= 0:
            continue
        expected_cells = [f"D{row}", f"E{row}", f"H{row}"]
        missing_cells = [cell for cell in expected_cells if _is_blank(sheet[cell].value)]
        _emit_story10b_table_block_finding(
            findings=findings,
            sheet_name="SUMMARY",
            block_id=f"SUMMARY.BLOCK.FUNDING_BY_DISEASE.R{row}",
            severity="warn",
            missing_cells=missing_cells,
            expected_cells=expected_cells,
        )

    contact_cells = ["B72", "C72", "E72", "G72", "H72"]
    _emit_story10b_table_block_finding(
        findings=findings,
        sheet_name="SUMMARY",
        block_id="SUMMARY.BLOCK.CONTACT_METADATA",
        severity="warn",
        missing_cells=[cell for cell in contact_cells if _is_blank(sheet[cell].value)],
        expected_cells=contact_cells,
    )

    authorization_cells = ["D84", "H85"]
    _emit_story10b_table_block_finding(
        findings=findings,
        sheet_name="SUMMARY",
        block_id="SUMMARY.BLOCK.AUTHORIZATION",
        severity="warn",
        missing_cells=[cell for cell in authorization_cells if _is_blank(sheet[cell].value)],
        expected_cells=authorization_cells,
    )


def _evaluate_story10b_shipment_table_blocks(*, sheet: Worksheet, findings: list[Finding]) -> None:
    consignee_sets = [
        (
            "SHIPMENT.BLOCK.CONSIGNEE_SET_1",
            "B8",
            ["C9", "C11", "C12", "C14", "C15", "F9", "F11", "F12", "F14", "F15"],
        ),
        (
            "SHIPMENT.BLOCK.CONSIGNEE_SET_2",
            "B17",
            ["C18", "C20", "C21", "C23", "C24", "F18", "F20", "F21", "F23", "F24"],
        ),
        (
            "SHIPMENT.BLOCK.CONSIGNEE_SET_3",
            "B26",
            ["C27", "C29", "C30", "C32", "C33", "F27", "F29", "F30", "F32", "F33"],
        ),
    ]
    for block_id, trigger_cell, expected_cells in consignee_sets:
        if _is_blank(sheet[trigger_cell].value):
            continue
        _emit_story10b_table_block_finding(
            findings=findings,
            sheet_name="SHIPMENT",
            block_id=block_id,
            severity="error",
            missing_cells=[cell for cell in expected_cells if _is_blank(sheet[cell].value)],
            expected_cells=expected_cells,
        )

    import_requirement_cells = ["H38", "H40", "H42", "H43"]
    _emit_story10b_table_block_finding(
        findings=findings,
        sheet_name="SHIPMENT",
        block_id="SHIPMENT.BLOCK.IMPORT_REQUIREMENTS",
        severity="error",
        missing_cells=[cell for cell in import_requirement_cells if _is_blank(sheet[cell].value)],
        expected_cells=import_requirement_cells,
    )

    if _is_yes(sheet["H38"].value):
        expected_cells = ["G39", "H39"]
        missing_cells = []
        if _to_float(sheet["G39"].value) is None:
            missing_cells.append("G39")
        if _is_blank(sheet["H39"].value):
            missing_cells.append("H39")
        _emit_story10b_table_block_finding(
            findings=findings,
            sheet_name="SHIPMENT",
            block_id="SHIPMENT.BLOCK.IMPORT_PERMIT_LEAD_TIME",
            severity="error",
            missing_cells=missing_cells,
            expected_cells=expected_cells,
        )

    if any(_is_yes(sheet[cell].value) for cell in import_requirement_cells):
        _emit_story10b_table_block_finding(
            findings=findings,
            sheet_name="SHIPMENT",
            block_id="SHIPMENT.BLOCK.IMPORT_DOCUMENTS",
            severity="warn",
            missing_cells=["B44"] if _is_blank(sheet["B44"].value) else [],
            expected_cells=["B44"],
        )


def _evaluate_story10b_summary_shipment_detailed(
    *,
    workbook: Workbook,
    findings: list[Finding],
    template_language: str | None,
) -> None:
    if not SUMMARY_SHIPMENT_REFERENCE_DOC_PATH.exists():
        return

    _build_finding(
        findings,
        rule_id="JRSM.SUMMARY_SHIPMENT.REFERENCE_APPLIED",
        severity="info",
        message="Detailed SUMMARY/SHIPMENT reference is configured; Story 10b deterministic checks applied.",
        recommendation="No action required.",
        expected=str(SUMMARY_SHIPMENT_REFERENCE_DOC_PATH),
        actual="REFERENCE_LOADED",
    )

    formula_specs: dict[str, dict[str, str]] = {
        "SUMMARY": STORY10B_SUMMARY_CANONICAL_FORMULAS,
        "SHIPMENT": STORY10B_SHIPMENT_CANONICAL_FORMULAS,
    }
    for sheet_name, canonical_by_cell in formula_specs.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        for cell_address, expected_formula in canonical_by_cell.items():
            actual_value = sheet[cell_address].value
            severity = _story10b_formula_severity(sheet_name, cell_address)
            localized_expected_formula = _localize_formula_literals(
                expected_formula,
                template_language=template_language,
            )
            if not _is_formula_value(actual_value):
                if _should_suppress_story10b_formula_finding(
                    findings=findings,
                    sheet_name=sheet_name,
                    cell_address=cell_address,
                    check_type="required",
                ):
                    continue
                _build_finding(
                    findings,
                    rule_id=f"JRSM.{sheet_name}.FORMULA_REQUIRED",
                    severity=severity,
                    message=f"{sheet_name} canonical formula cell is missing a formula.",
                    recommendation="Restore the canonical formula from template baseline for this cell.",
                    sheet=sheet_name,
                    cell=cell_address,
                    expected=localized_expected_formula,
                    actual=_format_actual_value(actual_value),
                )
                continue

            actual_formula = str(actual_value)
            if _normalize_formula(actual_formula) == _normalize_formula(localized_expected_formula):
                continue
            if _should_suppress_story10b_formula_finding(
                findings=findings,
                sheet_name=sheet_name,
                cell_address=cell_address,
                check_type="match",
            ):
                continue

            _build_finding(
                findings,
                rule_id=f"JRSM.{sheet_name}.FORMULA_MATCH",
                severity=severity,
                message=f"{sheet_name} canonical formula cell differs from expected formula.",
                recommendation="Restore the canonical formula from template baseline for this cell.",
                sheet=sheet_name,
                cell=cell_address,
                expected=localized_expected_formula,
                actual=actual_formula,
            )

    if "SUMMARY" in workbook.sheetnames:
        _evaluate_story10b_summary_table_blocks(sheet=workbook["SUMMARY"], findings=findings)
    if "SHIPMENT" in workbook.sheetnames:
        _evaluate_story10b_shipment_table_blocks(sheet=workbook["SHIPMENT"], findings=findings)


def validate_jrsm_workbook(
    *,
    workbook_path: Path,
    country: str,
    year_for_request_of_medicine: int,
    form_type: str = "jrsm",
    form_variation: str = "generic",
) -> ValidationEngineResult:
    """Run Story 4 workbook parsing + template conformity checks."""
    _ = (country, year_for_request_of_medicine, form_type, form_variation)
    findings: list[Finding] = []
    context: dict[str, Any] = {}

    try:
        workbook = load_workbook(filename=workbook_path, data_only=False)
    except Exception as exc:
        _build_finding(
            findings,
            rule_id="JRSM.WORKBOOK.FILE_PARSE",
            severity="error",
            message="Workbook parse failed.",
            recommendation="Upload a valid JRSM .xlsx or .xlsm workbook.",
            expected="Loadable workbook in formula-preserving mode (data_only=False).",
            actual=f"WORKBOOK_PARSE_FAILED: {exc}",
        )
        return ValidationEngineResult(
            status="completed",
            validation_outcome="fail",
            findings=findings,
            context=context,
        )

    missing_sheets = [sheet_name for sheet_name in REQUIRED_SHEETS if sheet_name not in workbook.sheetnames]
    if missing_sheets:
        missing_text = ", ".join(missing_sheets)
        _build_finding(
            findings,
            rule_id="JRSM.WORKBOOK.REQUIRED_SHEETS",
            severity="error",
            message=f"Workbook is missing required sheets: {missing_text}.",
            recommendation="Regenerate workbook from the JRSM template to restore required sheets.",
            expected=", ".join(REQUIRED_SHEETS),
            actual=missing_text,
        )
        _build_finding(
            findings,
            rule_id="JRSM.TEMPLATE.SHEET_SET",
            severity="warn",
            message="Workbook sheet set does not match expected JRSM template sheet set.",
            recommendation="Use a standard JRSM workbook generated from INTRO > Generate new form.",
            expected=", ".join(REQUIRED_SHEETS),
            actual=", ".join(workbook.sheetnames),
        )
    else:
        _build_finding(
            findings,
            rule_id="JRSM.TEMPLATE.SHEET_SET",
            severity="info",
            message="Workbook sheet set matches expected JRSM template sheet names.",
            recommendation="No action required.",
            expected=", ".join(REQUIRED_SHEETS),
            actual=", ".join(workbook.sheetnames),
        )

    intro_sheet = workbook["INTRO"] if "INTRO" in workbook.sheetnames else None
    template_profile = _detect_template_profile(workbook, intro_sheet)
    template_language = template_profile.get("language")
    template_supported = template_profile["status"] == SUPPORTED_TEMPLATE_STATUS
    _emit_template_profile_finding(findings, template_profile)
    if intro_sheet is None:
        _build_finding(
            findings,
            rule_id="JRSM.TEMPLATE.REQUIRED_CELLS",
            severity="error",
            message="Cannot validate required INTRO cells because INTRO sheet is missing.",
            recommendation="Provide a workbook containing the INTRO sheet.",
            sheet="INTRO",
            expected=", ".join([f"INTRO!{address}" for address in REQUIRED_INTRO_CELLS]),
            actual="INTRO sheet missing",
        )
        _build_finding(
            findings,
            rule_id=VERSION_MARKER_RULE_ID,
            severity="warn",
            message="Cannot evaluate template version marker because INTRO sheet is missing.",
            recommendation="Provide a workbook with INTRO!B2 populated by template generation.",
            sheet="INTRO",
            cell="B2",
            expected=EXPECTED_VERSION_MARKER,
            actual="INTRO sheet missing",
        )
    else:
        missing_cells = [cell_address for cell_address in REQUIRED_INTRO_CELLS if not _has_cell_address(intro_sheet, cell_address)]
        if missing_cells:
            _build_finding(
                findings,
                rule_id="JRSM.TEMPLATE.REQUIRED_CELLS",
                severity="error",
                message=f"Required INTRO cells are not addressable: {', '.join(missing_cells)}.",
                recommendation="Use a standard JRSM template workbook with required template cell layout.",
                sheet="INTRO",
                expected=", ".join([f"INTRO!{address}" for address in REQUIRED_INTRO_CELLS]),
                actual=", ".join([f"INTRO!{address}" for address in missing_cells]),
            )

        b2_value = intro_sheet["B2"].value if _has_cell_address(intro_sheet, "B2") else None
        b2_text = str(b2_value).strip() if b2_value is not None else ""
        if template_supported and template_language:
            _build_finding(
                findings,
                rule_id=VERSION_MARKER_RULE_ID,
                severity="info",
                message=(
                    f"Template version marker matches the supported "
                    f"{JRSM_LANGUAGE_PROFILES[template_language]['label']} JRSM signature."
                ),
                recommendation="No action required.",
                sheet="INTRO",
                cell="B2",
                expected=JRSM_LANGUAGE_PROFILES[template_language]["supported_markers"][0],
                actual=b2_text,
            )
        elif template_profile["status"] == UNSUPPORTED_VERSION_STATUS and template_language:
            _build_finding(
                findings,
                rule_id=VERSION_MARKER_RULE_ID,
                severity="warn",
                message=(
                    f"Template version marker does not match the supported "
                    f"{JRSM_LANGUAGE_PROFILES[template_language]['label']} JRSM signature."
                ),
                recommendation=(
                    f"Use the current supported {JRSM_LANGUAGE_PROFILES[template_language]['label']} "
                    "JRSM template lineage before relying on formula-governance findings."
                ),
                sheet="INTRO",
                cell="B2",
                expected=JRSM_LANGUAGE_PROFILES[template_language]["supported_markers"][0],
                actual=b2_text or "BLANK_OR_MISSING",
            )
        else:
            _build_finding(
                findings,
                rule_id=VERSION_MARKER_RULE_ID,
                severity="warn",
                message="Template version marker is missing or unrecognized for the supported JRSM language profiles.",
                recommendation="Confirm workbook was generated from the current supported English, French, or Spanish JRSM template lineage.",
                sheet="INTRO",
                cell="B2",
                expected=EXPECTED_VERSION_MARKER,
                actual=b2_text or "BLANK_OR_MISSING",
            )

        # Story 5: Request consistency checks against INTRO.
        e33_value = intro_sheet["E33"].value if _has_cell_address(intro_sheet, "E33") else None
        if _normalize_text(country) != _normalize_text(e33_value):
            _build_finding(
                findings,
                rule_id="JRSM.INTRO.COUNTRY_MATCH",
                severity="error",
                message="Country mismatch between request and workbook.",
                recommendation="Correct request country or update INTRO!E33 to the expected country name.",
                sheet="INTRO",
                cell="E33",
                expected=country,
                actual=str(e33_value).strip() if e33_value is not None else "BLANK_OR_MISSING",
            )

        e35_value = intro_sheet["E35"].value if _has_cell_address(intro_sheet, "E35") else None
        request_year = _normalize_year(year_for_request_of_medicine)
        workbook_year = _normalize_year(e35_value)
        if request_year != workbook_year:
            _build_finding(
                findings,
                rule_id="JRSM.INTRO.YEAR_MATCH",
                severity="error",
                message="Request year does not match INTRO!E35.",
                recommendation="Correct request year or update INTRO!E35 so they match numerically.",
                sheet="INTRO",
                cell="E35",
                expected=str(request_year) if request_year is not None else str(year_for_request_of_medicine),
                actual=str(workbook_year) if workbook_year is not None else str(e35_value).strip() or "BLANK_OR_MISSING",
            )

        # Story 5 required-value checks apply to driver input cells (not spacer cells).
        for cell_address in REQUIRED_INTRO_VALUE_CELLS:
            cell_value = intro_sheet[cell_address].value if _has_cell_address(intro_sheet, cell_address) else None
            if _is_blank(cell_value):
                _build_finding(
                    findings,
                    rule_id="JRSM.INTRO.REQUIRED_FIELDS_PRESENT",
                    severity="error",
                    message="Required INTRO field is blank.",
                    recommendation="Populate this required INTRO field and rerun validation.",
                    sheet="INTRO",
                    cell=cell_address,
                    expected="Non-empty value",
                    actual="BLANK_OR_MISSING",
                )

        # Story 6: Endemicity/visibility derivation + IU row-window engine.
        e37_value = intro_sheet["E37"].value if _has_cell_address(intro_sheet, "E37") else None
        e39_value = intro_sheet["E39"].value if _has_cell_address(intro_sheet, "E39") else None
        e41_value = intro_sheet["E41"].value if _has_cell_address(intro_sheet, "E41") else None
        e43_value = intro_sheet["E43"].value if _has_cell_address(intro_sheet, "E43") else None
        has_lf = not _matches_template_semantic_literal(
            e37_value,
            template_language=template_language,
            semantic_key="non_endemic",
        )
        has_oncho = not _matches_template_semantic_literal(
            e39_value,
            template_language=template_language,
            semantic_key="non_endemic",
        )
        has_sth = not _matches_template_semantic_literal(
            e41_value,
            template_language=template_language,
            semantic_key="non_endemic",
        )
        has_sch = not _matches_template_semantic_literal(
            e43_value,
            template_language=template_language,
            semantic_key="non_endemic",
        )

        n_raw_value = intro_sheet["E45"].value if _has_cell_address(intro_sheet, "E45") else None
        n_value = _normalize_positive_int(n_raw_value)
        iu_start_row = 10
        iu_end_row: int | None = None
        if n_value is None:
            _build_finding(
                findings,
                rule_id="JRSM.IU_WINDOW.N_VALID",
                severity="error",
                message="INTRO!E45 is not a valid positive integer for active IU count.",
                recommendation="Populate INTRO!E45 with a positive integer number of administrative units.",
                sheet="INTRO",
                cell="E45",
                expected="Positive integer",
                actual=str(n_raw_value).strip() if n_raw_value is not None else "BLANK_OR_MISSING",
            )
            _build_finding(
                findings,
                rule_id="JRSM.IU_WINDOW.RANGE_ENFORCED",
                severity="warn",
                message="Active IU row window could not be derived because N is invalid.",
                recommendation="Correct INTRO!E45, then rerun validation to derive active IU row range.",
                expected="Rows 10..(9+N)",
                actual="UNAVAILABLE",
            )
        else:
            iu_end_row = 9 + n_value
            _build_finding(
                findings,
                rule_id="JRSM.IU_WINDOW.N_VALID",
                severity="info",
                message="INTRO!E45 contains a valid IU row count.",
                recommendation="No action required.",
                sheet="INTRO",
                cell="E45",
                expected="Positive integer",
                actual=str(n_value),
            )
            _build_finding(
                findings,
                rule_id="JRSM.IU_WINDOW.RANGE_ENFORCED",
                severity="info",
                message="Active IU row window derived from INTRO!E45.",
                recommendation="Use this IU window for row-level deterministic checks.",
                expected="Rows 10..(9+N)",
                actual=f"Rows {iu_start_row}..{iu_end_row}",
            )

        has_ida = False
        has_ida_known = False
        country_info_sheet = workbook["COUNTRY_INFO"] if "COUNTRY_INFO" in workbook.sheetnames else None
        if n_value is None:
            _build_finding(
                findings,
                rule_id="JRSM.ENDEMICITY.IDA_DETECTION",
                severity="warn",
                message="IDA detection skipped because active IU range is unavailable.",
                recommendation="Correct INTRO!E45 so IDA detection can scan COUNTRY_INFO column Z.",
                expected="Any non-empty value in COUNTRY_INFO!Z10:Z(9+N)",
                actual="NOT_EVALUATED",
            )
        elif country_info_sheet is None:
            _build_finding(
                findings,
                rule_id="JRSM.ENDEMICITY.IDA_DETECTION",
                severity="warn",
                message="IDA detection skipped because COUNTRY_INFO sheet is missing.",
                recommendation="Provide a workbook with COUNTRY_INFO sheet present.",
                expected="COUNTRY_INFO!Z10:Z(9+N) available",
                actual="COUNTRY_INFO missing",
            )
        else:
            has_ida_known = True
            assert iu_end_row is not None
            for row in range(iu_start_row, iu_end_row + 1):
                z_value = country_info_sheet[f"Z{row}"].value
                if not _is_blank(z_value):
                    has_ida = True
                    break
            _build_finding(
                findings,
                rule_id="JRSM.ENDEMICITY.IDA_DETECTION",
                severity="info",
                message="IDA criteria detection evaluated from COUNTRY_INFO active IU rows.",
                recommendation="No action required.",
                expected=f"Scan COUNTRY_INFO!Z{iu_start_row}:Z{iu_end_row} for any non-empty value",
                actual="IDA_DETECTED" if has_ida else "IDA_NOT_DETECTED",
            )

        expected_sheet_visibility = {sheet_name: True for sheet_name in CORE_EXPECTED_VISIBLE_SHEETS}
        expected_sheet_visibility.update(
            {
                "DEC": has_lf and not has_oncho,
                "IVM": has_oncho,
                "IVM+": has_ida if has_ida_known else False,
                "ALB_MBD": has_lf or has_sth,
                "PZQ": has_sch,
            }
        )
        reviewable_sheet_visibility = _build_reviewable_sheet_visibility(
            workbook=workbook,
            expected_sheet_visibility=expected_sheet_visibility,
        )
        actual_sheet_visibility_text = ", ".join(
            [
                f"{sheet}={workbook[sheet].sheet_state}"
                for sheet in MODULE_VISIBILITY_SHEETS
                if sheet in workbook.sheetnames
            ]
        )
        expected_visible_module_count = sum(1 for sheet_name in MODULE_VISIBILITY_SHEETS if expected_sheet_visibility[sheet_name])
        _build_finding(
            findings,
            rule_id="JRSM.ENDEMICITY.EXPECTED_MODULE_VISIBILITY",
            severity="info",
            message="Expected module visibility profile was derived from INTRO endemicity and IDA signals.",
            recommendation="Use this profile as the baseline for sheet-state mismatch checks.",
            expected=", ".join(
                [f"{sheet}={'visible' if expected_sheet_visibility[sheet] else 'hidden'}" for sheet in MODULE_VISIBILITY_SHEETS]
            ),
            actual=actual_sheet_visibility_text,
        )

        for sheet_name, expected_visible in expected_sheet_visibility.items():
            if sheet_name not in workbook.sheetnames:
                continue
            actual_state = workbook[sheet_name].sheet_state
            actual_visible = actual_state == "visible"
            if actual_visible == expected_visible:
                continue

            _build_finding(
                findings,
                rule_id="JRSM.ENDEMICITY.ACTUAL_SHEET_STATE_MISMATCH",
                severity="warn",
                message=f"Sheet visibility mismatch for {sheet_name}.",
                recommendation="Review sheet visibility against expected endemicity-driven profile.",
                sheet=sheet_name,
                expected="visible" if expected_visible else "hidden",
                actual=actual_state,
            )

        context = {
            "template_language": template_language,
            "template_profile_status": template_profile["status"],
            "has_lf": has_lf,
            "has_oncho": has_oncho,
            "has_sth": has_sth,
            "has_sch": has_sch,
            "has_ida": has_ida if has_ida_known else None,
            "visibility_driver_flags": {
                "has_lf": has_lf,
                "has_oncho": has_oncho,
                "has_sth": has_sth,
                "has_sch": has_sch,
                "has_ida": has_ida if has_ida_known else None,
            },
            "n": n_value,
            "iu_window_start_row": iu_start_row if n_value is not None else None,
            "iu_window_end_row": iu_end_row,
            "expected_visible_module_count": expected_visible_module_count,
            "expected_sheet_visibility": expected_sheet_visibility,
            "reviewable_sheet_visibility": reviewable_sheet_visibility,
        }

        # Story 7: COUNTRY_INFO active-row completeness checks.
        if country_info_sheet is not None and iu_end_row is not None:
            active_rows = list(range(iu_start_row, iu_end_row + 1))

            # Required country ID columns B:D for every active IU row.
            for row in active_rows:
                for column in ["B", "C", "D"]:
                    cell_address = f"{column}{row}"
                    if _is_blank(country_info_sheet[cell_address].value):
                        _build_finding(
                            findings,
                            rule_id="JRSM.COUNTRY_INFO.REQUIRED_ID_COLUMNS_B_TO_D",
                            severity="error",
                            message="COUNTRY_INFO active IU row is missing required ID field.",
                            recommendation="Populate COUNTRY_INFO columns B, C, and D for each active IU row.",
                            sheet="COUNTRY_INFO",
                            cell=cell_address,
                            expected="Non-empty value",
                            actual="BLANK_OR_MISSING",
                        )

            # Endemicity-driven H:K checks.
            endemicity_required_columns = {
                "H": has_lf,
                "I": has_oncho,
                "J": has_sth,
                "K": has_sch,
            }
            for column, is_required in endemicity_required_columns.items():
                if not is_required:
                    continue

                blank_count = 0
                for row in active_rows:
                    cell_address = f"{column}{row}"
                    if _is_blank(country_info_sheet[cell_address].value):
                        blank_count += 1
                        _build_finding(
                            findings,
                            rule_id="JRSM.COUNTRY_INFO.ENDEMICITY_COLUMNS_H_TO_K",
                            severity="error",
                            message="COUNTRY_INFO endemicity-required disease column is blank for active IU row.",
                            recommendation="Populate required endemicity code columns H, I, J, K for active rows based on endemicity profile.",
                            sheet="COUNTRY_INFO",
                            cell=cell_address,
                            expected="Non-empty value",
                            actual="BLANK_OR_MISSING",
                        )

                if blank_count == len(active_rows) and _header_text(country_info_sheet, column):
                    _build_finding(
                        findings,
                        rule_id="JRSM.COUNTRY_INFO.ENDEMICITY_COLUMNS_H_TO_K",
                        severity="error",
                        message=f"COUNTRY_INFO column {column} is systematically blank across active IU rows.",
                        recommendation="Populate the endemicity-required column for active IU rows or correct endemicity configuration.",
                        sheet="COUNTRY_INFO",
                        cell=f"{column}{iu_start_row}:{column}{iu_end_row}",
                        expected="At least one non-empty value in active IU window",
                        actual="SYSTEMATICALLY_BLANK",
                    )

            # Required output block V:AD when header-populated and configuration-expected.
            required_output_columns: list[str] = []
            for column_index in range(column_index_from_string("V"), column_index_from_string("AD") + 1):
                column = get_column_letter(column_index)
                if _is_column_hidden(country_info_sheet, column):
                    continue
                header = _header_text(country_info_sheet, column)
                if not _is_output_column_expected(
                    header_text=header,
                    has_lf=has_lf,
                    has_oncho=has_oncho,
                    has_sth=has_sth,
                    has_sch=has_sch,
                    has_ida=has_ida if has_ida_known else None,
                ):
                    continue
                if not _is_context_optional_country_info_output_column_required(
                    sheet=country_info_sheet,
                    column_letter=column,
                    active_rows=active_rows,
                ):
                    continue
                required_output_columns.append(column)

                blank_rows = [row for row in active_rows if _is_blank(country_info_sheet[f"{column}{row}"].value)]
                column_descriptor = f"column {column} ({_column_label(country_info_sheet, column)})"
                if len(blank_rows) == len(active_rows):
                    _build_finding(
                        findings,
                        rule_id="JRSM.COUNTRY_INFO.REQUIRED_OUTPUT_COLUMNS_V_TO_AD",
                        severity="error",
                        message=f"COUNTRY_INFO {column_descriptor} is blank across active IU rows {iu_start_row}:{iu_end_row}.",
                        recommendation="Review whether this header is intentionally blank for the current configuration; otherwise populate it and rerun validation.",
                        sheet="COUNTRY_INFO",
                        cell=f"{column}{iu_start_row}:{column}{iu_end_row}",
                        expected="At least one non-empty value in active IU window",
                        actual="SYSTEMATICALLY_BLANK",
                    )
                    continue

                for row in blank_rows:
                    cell_address = f"{column}{row}"
                    _build_finding(
                        findings,
                        rule_id="JRSM.COUNTRY_INFO.REQUIRED_OUTPUT_COLUMNS_V_TO_AD",
                        severity="error",
                        message=f"COUNTRY_INFO {column_descriptor} is blank for active IU row {row}.",
                        recommendation="Populate required COUNTRY_INFO output cells for active IU rows when the column is expected for the current configuration.",
                        sheet="COUNTRY_INFO",
                        cell=cell_address,
                        expected="Non-empty value",
                        actual="BLANK_OR_MISSING",
                    )

            context["country_info_required_output_columns"] = required_output_columns

        # Story 8: medicine module row completeness checks on active IU rows.
        if iu_end_row is not None:
            active_rows = list(range(iu_start_row, iu_end_row + 1))
            for (
                module_sheet_name,
                required_columns,
                required_rule_id,
                _,
            ) in MODULE_COMPLETENESS_SPECS:
                if module_sheet_name not in workbook.sheetnames:
                    continue

                module_sheet = workbook[module_sheet_name]
                if not reviewable_sheet_visibility.get(module_sheet_name, False):
                    continue

                reviewable_required_columns = [
                    column for column in required_columns if not _is_column_hidden(module_sheet, column)
                ]
                if not reviewable_required_columns:
                    continue

                blank_rows_by_column = {
                    column: [row for row in active_rows if _is_blank(module_sheet[f"{column}{row}"].value)]
                    for column in reviewable_required_columns
                }
                all_blank_columns = [
                    column for column in reviewable_required_columns if len(blank_rows_by_column[column]) == len(active_rows)
                ]
                for grouped_columns in _group_contiguous_columns(all_blank_columns):
                    _build_finding(
                        findings,
                        rule_id=required_rule_id,
                        severity="error",
                        message=(
                            f"{module_sheet_name} {_format_column_descriptor(module_sheet, grouped_columns)} "
                            f"are blank across active IU rows {iu_start_row}:{iu_end_row}."
                            if len(grouped_columns) > 1
                            else f"{module_sheet_name} {_format_column_descriptor(module_sheet, grouped_columns)} "
                            f"is blank across active IU rows {iu_start_row}:{iu_end_row}."
                        ),
                        recommendation=f"Review whether these required module columns are intentionally blank; otherwise populate them and rerun validation for {module_sheet_name}.",
                        sheet=module_sheet_name,
                        cell=_format_grouped_cell_range(grouped_columns, iu_start_row, iu_end_row),
                        expected="At least one non-empty value in active IU window",
                        actual="SYSTEMATICALLY_BLANK",
                    )

                for column in reviewable_required_columns:
                    if column in all_blank_columns:
                        continue
                    column_descriptor = f"column {column} ({_column_label(module_sheet, column)})"
                    for row in blank_rows_by_column[column]:
                        cell_address = f"{column}{row}"
                        _build_finding(
                            findings,
                            rule_id=required_rule_id,
                            severity="error",
                            message=f"{module_sheet_name} {column_descriptor} is blank for active IU row {row}.",
                            recommendation=f"Populate required module cells in {module_sheet_name} and rerun validation.",
                            sheet=module_sheet_name,
                            cell=cell_address,
                            expected="Non-empty value",
                            actual="BLANK_OR_MISSING",
                        )

        # Story 9: submitter-day formula governance checks.
        if template_supported:
            _evaluate_story9_formula_governance(
                workbook=workbook,
                findings=findings,
                iu_start_row=iu_start_row,
                iu_end_row=iu_end_row,
                template_language=template_language,
                reviewable_sheet_visibility=reviewable_sheet_visibility,
            )
        _evaluate_story10_summary_shipment_placeholder(
            workbook=workbook,
            findings=findings,
        )
        if template_supported:
            _evaluate_story10b_summary_shipment_detailed(
                workbook=workbook,
                findings=findings,
                template_language=template_language,
            )

    return ValidationEngineResult(
        status="completed",
        validation_outcome=_derive_outcome(findings),
        findings=findings,
        context=context,
    )
